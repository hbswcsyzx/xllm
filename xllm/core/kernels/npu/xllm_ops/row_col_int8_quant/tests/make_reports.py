# Copyright 2026 The xLLM Authors. SPDX-License-Identifier: Apache-2.0
"""Populate the official seven-file report template from saved NPU evidence."""

import argparse
import json
import zipfile
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.styles import Alignment


def _load(path: Path) -> Any:
    return json.loads(path.read_text())


def _workbook(template: Path, name: str, summary: str) -> Any:
    book = load_workbook(template / name)
    sheet = book.active
    # The published template contains unrelated example screenshots/results.
    sheet._images.clear()
    for row in sheet.iter_rows(min_row=4):
        for cell in row:
            cell.value = None
    sheet["A1"] = summary
    sheet["A1"].alignment = Alignment(wrap_text=True, vertical="top")
    return book


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--task-root", type=Path, required=True)
    parser.add_argument("--template", type=Path, required=True)
    parser.add_argument("--design-pr-url", default="")
    args = parser.parse_args()
    task = args.task_root.resolve()
    op = Path(__file__).resolve().parents[1]
    destination = op / "task_submission"
    destination.mkdir(exist_ok=True)
    public = _load(task / "docs/public-final.json")
    engineering = _load(task / "docs/engineering-final.json")
    performance = _load(task / "docs/performance-final.json")
    acceptance = _load(task / "docs/performance-acceptance.json")
    baselines = _load(task / "docs/baselines-final.json")
    cases = _load(task / "material/public/case.json")
    by_name = {case["case_name"]: case for case in cases}
    assert all(report["status"] == "passed" for report in (public, engineering, performance, acceptance, baselines))
    assert len(public["cases"]) == 27 and len(engineering["records"]) == 115
    assert len(performance["cases"]) == 5 and len(acceptance["checks"]) == 5
    environment = "Ascend910_9362 / 40 AIV / driver 25.5.5 / CANN 9.1.0 / Python 3.12.13 / torch 2.10.0+cpu / torch_npu 2.10.0.post4+git5dd8ef3"
    accuracy_name = "2.1 精度自验证报告.xlsx"
    accuracy = _workbook(
        args.template,
        accuracy_name,
        f"本地NPU精度通过。27个公开+12个额外偏移输入；115条扩展工程记录均通过。\n{environment}\nINT8仅允许任务规定的半整数邻域差异；精确中点/全零严格一致，FP32 absmax逐位一致。实际设备与任务指定9382/48 AIV不同，外部验收未执行。",
    )
    sheet = accuracy.active
    sheet["D3"] = "日志证据"
    entries = [(item["case_name"], item["outputs"]) for item in public["cases"]]
    entries += [
        (item["case_name"], item["checks"])
        for item in engineering["records"]
        if item["kind"] == "extra" and item["entry"] == "run"
    ]
    assert len(entries) == 39
    for index, (name, checks) in enumerate(entries, 4):
        case = by_name[name.split("__extra_seed_")[0]]
        values = [
            str(case["input_desc"][0]["shape"]),
            "float16",
            "通过",
            "2.2 精度自验证日志.log",
            name + "\n" + json.dumps(checks, ensure_ascii=False),
        ]
        for column, value in enumerate(values, 1):
            sheet.cell(index, column, value).alignment = Alignment(wrap_text=True)
    extra = accuracy.create_sheet("工程门禁")
    extra.append(["kind", "case_name", "entry", "status", "details"])
    for item in engineering["records"]:
        extra.append(
            [item["kind"], item["case_name"], item["entry"], item["status"], json.dumps(item, ensure_ascii=False)]
        )
    accuracy.save(destination / accuracy_name)

    performance_name = "3.1 性能自验证报告.xlsx"
    timing = _workbook(
        args.template,
        performance_name,
        f"5/5固定耗时门槛通过；正式run每次新输出。50次预热，5轮×200次，轮均墙钟中位数，保留全部轮次。\n{environment}\n标杆在同机按相同协议重测；设备Event包含提交间隙，profiler时间不作为验收时间。",
    )
    sheet = timing.active
    sheet["D3"] = "日志证据"
    raw = timing.create_sheet("原始轮次")
    raw.append(["case", "implementation", "round", "wall_us", "event_us"])
    baseline_map = {item["case_name"]: item for item in baselines["cases"]}
    limits = {item["case_name"]: item["max_wall_us"] for item in acceptance["checks"]}
    for index, item in enumerate(performance["cases"], 4):
        name = item["case_name"]
        custom = item["results"]["custom"]
        baseline = baseline_map[name]
        selected = baseline["candidates"][baseline["selected"]]
        values = [
            str(by_name[name]["input_desc"][0]["shape"]),
            "float16",
            "通过",
            "3.2 性能自验证日志.log",
            selected["wall_us_median"],
            custom["wall_us_median"],
            custom["wall_us_median"] / selected["wall_us_median"],
            f"{name};固定上限={limits[name]} us;标杆={baseline['selected']}",
        ]
        for column, value in enumerate(values, 1):
            sheet.cell(index, column, value).alignment = Alignment(wrap_text=True)
        for label, result in [("custom", custom)] + [
            (name, value) for name, value in baseline["candidates"].items() if value["status"] == "qualified"
        ]:
            for round_index, (wall, event) in enumerate(
                zip(result["wall_us_rounds"], result["device_event_us_rounds"]), 1
            ):
                raw.append([name, label, round_index, wall, event])
    timing.save(destination / performance_name)

    memory_name = "4.1 内存自验证报告.xlsx"
    memory = _workbook(
        args.template,
        memory_name,
        f"5个性能输入×ACLNN/直调两入口，10条内存检查全部通过。峰值为预热后活跃张量增量；每入口重复20次并同步后回到调用前。\n{environment}\n绝对上限=输出字节+任务workspace上限+16 KiB；无全矩阵FP32/转置workspace。",
    )
    sheet = memory.active
    sheet["D3"] = "日志证据"
    sheet["E3"] = "本题允许峰值增量(GB)"
    sheet["F3"] = "实际峰值增量(GB)"
    sheet["G3"] = "实际/上限"
    for index, item in enumerate((item for item in engineering["records"] if item["kind"] == "memory"), 4):
        allowed = item["output_bytes"] + item["workspace_limit"] + 16 * 1024
        values = [
            str(by_name[item["case_name"]]["input_desc"][0]["shape"]),
            "float16",
            "通过",
            "4.2 内存自验证日志.log",
            allowed / 1e9,
            item["peak_increment"] / 1e9,
            item["peak_increment"] / allowed,
            json.dumps(item, ensure_ascii=False),
        ]
        for column, value in enumerate(values, 1):
            sheet.cell(index, column, value).alignment = Alignment(wrap_text=True)
    memory.save(destination / memory_name)

    logs = task / "docs/logs"
    accuracy_log = (
        (logs / "public-final.log").read_text()
        + "\n"
        + (logs / "engineering-final.log").read_text()
        + "\n"
        + (logs / "example-final.log").read_text()
    )
    (destination / "2.2 精度自验证日志.log").write_text(accuracy_log)
    performance_log = (
        (logs / "performance-final.log").read_text()
        + "\n"
        + (logs / "baselines-final.log").read_text()
        + "\n"
        + json.dumps(acceptance, ensure_ascii=False, indent=2)
    )
    (destination / "3.2 性能自验证日志.log").write_text(performance_log)
    (destination / "4.2 内存自验证日志.log").write_text(
        "\n".join(
            json.dumps(item, ensure_ascii=False)
            for item in engineering["records"]
            if item["kind"] in {"memory", "stream", "input_early_release", "output_lifetime"}
        )
        + "\n"
    )
    design_review = args.design_pr_url or "未提交"
    steps = f"""# RowColInt8Quant 自验证步骤

本地全部门禁通过；本报告不等同于指定9382硬件复验、设计评审或平台验收。

设计文档评审：{design_review}。PR提交与评审通过分别记录，提交不代表评审或平台验收通过。

报告、原始日志、安装包、预编译二进制及验收ZIP仅在本地交付；远端分支只提交源码、构建脚本、文档和测试代码。

环境：{environment}。任务指定9382/48 AIV、驱动25.5.t6.b001，与实测差异保留。

所有命令从任务worktree经 ../run.sh 执行；源码/构建只使用本题worktree及state，私有安装不覆盖系统CANN。

```bash
op=xllm/core/kernels/npu/xllm_ops/row_col_int8_quant
../run.sh bash "$op/build.sh"
../run.sh python3 "$op/build_bridge.py"
../run.sh bash "$op/examples/run.sh"
../run.sh python3 ../material/public/verify.py --solution-file "$op/solution.py" --report ../docs/public-final.json
../run.sh python3 "$op/tests/engineering.py" --public ../material/public --report ../docs/engineering-final.json
../run.sh python3 ../material/public/benchmark.py --solution-file "$op/solution.py" --report ../docs/performance-final.json
../run.sh python3 ../material/public/check_performance.py --report ../docs/performance-final.json --output ../docs/performance-acceptance.json
../run.sh python3 "$op/tests/baselines.py" --public ../material/public --report ../docs/baselines-final.json
```

官方case SHA256=f0e6fdbf77c27960762db89253044323244e7ab4b549d2bf4e4d4d6fe0d8be76。
官方benchmark/check_performance未修改。正式run包含新输出分配、ACLNN提交至设备完成；输入生成/传输、标杆计算和比较均在计时外。
50次预热、5轮每轮200次，全部wall/Event轮次保留，以轮均wall中位数对照32/35/73/244/42 us固定门槛。

27公开+12额外偏移输入共39个有效输入。13类非法场景在prepare/run/run_direct共39条；115条工程记录包括直调公开27、额外24、非法39、native元数据6、stream2、输入提前释放2、输出生命周期2、内存10、验证器3。
全部输入按原始FP16位模式传输；额外12个指定输入offset=1。每次输出独立连续、无输入修改，absmax逐位一致。内存按活跃张量统计；run记录ACLNN实际workspace查询值，run_direct记录同tiling计划的显式分配字节，分别核对同一上限。

功能必须覆盖完整合法尺寸域，而非仅五个性能shape。27个冻结公开用例全部必测，另有12个强制额外输入及工程门禁。rank严格使用(2,2,2)；非连续输入保持(257,1026)，仅stride改为(2052,2)。负向日志保留输入shape/stride/dtype/device以核对场景隔离。

完整原始JSON、选定PMU CSV、源码及私有.run/bridge/kernel指纹见交付manifest和docs。Git分支及固定commit见交付清单；平台设计评审/验收状态另列，不以本地PASS替代。
"""
    (destination / "1 自验证步骤说明.md").write_text(steps)
    names = [
        "1 自验证步骤说明.md",
        accuracy_name,
        "2.2 精度自验证日志.log",
        performance_name,
        "3.2 性能自验证日志.log",
        memory_name,
        "4.2 内存自验证日志.log",
    ]
    with zipfile.ZipFile(op / "task_submission.zip", "w", zipfile.ZIP_DEFLATED) as archive:
        for name in names:
            archive.write(destination / name, name)


if __name__ == "__main__":
    main()
